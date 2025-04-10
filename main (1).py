import os
import psycopg2
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from datetime import datetime
from dateutil.relativedelta import relativedelta
import sys
import copy

def main():
    # 1. Obtener credenciales y la ruta del archivo base
    db_name = os.environ.get('DB_NAME', 'semillas')
    db_user = os.environ.get('DB_USER', 'openerp')
    db_password = os.environ.get('DB_PASSWORD', '')
    db_host = os.environ.get('DB_HOST', '2.136.142.253')
    db_port = os.environ.get('DB_PORT', '5432')
    # Archivo a escribir
    file_path = os.environ.get('EXCEL_FILE_PATH', 'Listado de Facturas Open Trading-Sem 2025.xlsx')
    
    db_params = {
        'dbname': db_name,
        'user': db_user,
        'password': db_password,
        'host': db_host,
        'port': db_port
    }
    
    # 2. Calcular el rango de fechas dinámico:
    # Desde el primer día del mes de hace dos meses hasta el día actual.
    end_date = datetime.now()
    start_date = (end_date - relativedelta(months=2)).replace(day=1)
    end_date_str = end_date.strftime('%Y-%m-%d')
    start_date_str = start_date.strftime('%Y-%m-%d')
    
    # 3. Consulta SQL utilizando el rango dinámico de fechas
    query = f"""
    SELECT 
        rc.name as "Compañía",
        a.internal_number as "Número",
        rp.vat as "CIF",
        p.name AS "Cliente/Proveedor",
        CASE 
            WHEN a.type = 'out_invoice' THEN 'Cliente'
            WHEN a.type = 'in_invoice' THEN 'Proveedor'
            WHEN a.type = 'out_refund' THEN 'Rectificativa Cliente'
            WHEN a.type = 'in_refund' THEN 'Rectificativa Proveedor'
            ELSE 'Otro'
        END as "Tipo",
 
        to_char(a.date_invoice, 'DD') as "Fecha Dia",
        to_char(a.date_invoice, 'MM') as "Fecha Mes",
        to_char(a.date_invoice, 'YYYY') as "Fecha Año",
        to_char(a.date_invoice, 'DD/MM/YYYY') as "Fecha",
 
        afp.name as "Régimen fiscal",
 
        a.amount_untaxed_0 as "Base 0", 
        a.amount_taxed_0 as "Importe 0", 
        a.amount_total_0 as "Total 0",
 
        a.amount_untaxed_4 as "Base 4", 
        a.amount_taxed_4 as "Importe 4", 
        a.amount_total_4 as "Total 4",
 
        a.amount_untaxed_8 as "Base 8", 
        a.amount_taxed_8 as "Importe 8", 
        a.amount_total_8 as "Total 8",
 
        a.amount_untaxed_10 as "Base 10", 
        a.amount_taxed_10 as "Importe 10", 
        a.amount_total_10 as "Total 10",
 
        a.amount_untaxed_18 as "Base 18", 
        a.amount_taxed_18 as "Importe 18", 
        a.amount_total_18 as "Total 18",
 
        a.amount_untaxed_21 as "Base 21", 
        a.amount_taxed_21 as "Importe 21", 
        a.amount_total_21 as "Total 21",
 
        a.recargo_equivalencia_05 as "Recargo 0,5", 
        a.recargo_equivalencia_1 as "Recargo 1", 
        a.recargo_equivalencia_14 as "Recargo 1.4", 
        a.recargo_equivalencia_4 as "Recargo 4",
        a.recargo_equivalencia_52 as "Recargo 5.2",
 
        a.amount_total as "Total",
 
        (CASE 
            WHEN a.type = 'out_refund' THEN
                - (
                    COALESCE(a.amount_untaxed_0, 0) +
                    COALESCE(a.amount_untaxed_4, 0) +
                    COALESCE(a.amount_untaxed_8, 0) +
                    COALESCE(a.amount_untaxed_10, 0) +
                    COALESCE(a.amount_untaxed_18, 0) +
                    COALESCE(a.amount_untaxed_21, 0)
                )
            ELSE
                (
                    COALESCE(a.amount_untaxed_0, 0) +
                    COALESCE(a.amount_untaxed_4, 0) +
                    COALESCE(a.amount_untaxed_8, 0) +
                    COALESCE(a.amount_untaxed_10, 0) +
                    COALESCE(a.amount_untaxed_18, 0) +
                    COALESCE(a.amount_untaxed_21, 0)
                )
         END) as "Total BASES",
        CASE   
            WHEN a.obsolescencia = TRUE THEN 'SI' 
            ELSE 'NO' 
        END as "BULK",
        'S-' || rp.id AS "ID Cliente",
        rpa.city AS "Ciudad"
 
    FROM 
        account_invoice as a 
    INNER JOIN 
        res_partner as p ON p.id = a.partner_id 
    INNER JOIN 
        account_fiscal_position afp ON afp.id = a.fiscal_position 
    INNER JOIN 
        res_company rc ON rc.id = a.company_id
    INNER JOIN 
        res_partner rp ON rp.id = a.partner_id
    INNER JOIN 
        res_partner_address rpa ON rpa.id = a.address_invoice_id
 
    WHERE 
        a.state IN ('open','paid') 
        AND a.date_invoice BETWEEN '{start_date_str}' AND '{end_date_str}'
        AND a.type IN ('out_invoice', 'out_refund')
        AND a.internal_number NOT LIKE 'RA%'
 
    ORDER BY 
        a.internal_number;
    """
    
    # 4. Conectar a la base de datos, ejecutar la consulta y obtener los datos
    try:
        with psycopg2.connect(**db_params) as conn:
            with conn.cursor() as cur:
                cur.execute(query)
                resultados = cur.fetchall()
                headers = [desc[0] for desc in cur.description]
    except Exception as e:
        print(f"Error al conectar o ejecutar la consulta: {e}")
        sys.exit(1)
    
    if not resultados:
        print("No se obtuvieron resultados de la consulta.")
        return
    else:
        print(f"Se obtuvieron {len(resultados)} filas de la consulta.")
    
    # 5. Abrir el archivo base: Listado de Facturas Open Trading-Sem 2025.xlsx
    try:
        book = load_workbook(file_path)
        sheet = book.active
    except FileNotFoundError:
        print(f"No se encontró el archivo base '{file_path}'. Se aborta para no perder el formato.")
        return
    
    # 6. Evitar duplicados usando la columna "Número" (segunda columna del resultado)
    existing_numbers = {row[1] for row in sheet.iter_rows(min_row=2, values_only=True)}
    for row in resultados:
        if row[1] not in existing_numbers:
            sheet.append(row)
            new_row_index = sheet.max_row
            # Copiar el formato de la fila anterior para mantener consistencia
            if new_row_index > 1:
                for col in range(1, sheet.max_column + 1):
                    source_cell = sheet.cell(row=new_row_index - 1, column=col)
                    target_cell = sheet.cell(row=new_row_index, column=col)
                    target_cell.font = copy.copy(source_cell.font)
                    target_cell.fill = copy.copy(source_cell.fill)
                    target_cell.border = copy.copy(source_cell.border)
                    target_cell.alignment = copy.copy(source_cell.alignment)
    
    # 7. Actualizar la referencia de la tabla existente (se asume que la tabla se llama "Portes")
    if "Portes" in sheet.tables:
        tabla = sheet.tables["Portes"]
        max_row = sheet.max_row
        max_col = sheet.max_column
        last_col_letter = get_column_letter(max_col)
        new_ref = f"A1:{last_col_letter}{max_row}"
        tabla.ref = new_ref
        print(f"Tabla 'Portes' actualizada a rango: {new_ref}")
    else:
        print("No se encontró la tabla 'Portes'. Se conservará el formato actual, pero no se actualizará la referencia de la tabla.")
    
    # 8. Guardar el archivo Excel actualizado
    book.save(file_path)
    print(f"Archivo guardado con la estructura de tabla en '{file_path}'.")

if __name__ == '__main__':
    main()
